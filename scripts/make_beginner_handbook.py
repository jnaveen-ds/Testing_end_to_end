"""Generate the beginner handbook's HTML and Word formats from its Markdown source.

Install the small documentation-only dependencies first when needed:
    python -m pip install -r scripts/requirements-docs.txt
"""

from __future__ import annotations

import re
from pathlib import Path

import markdown
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "AZURE_GENAI_BEGINNER_HANDBOOK.md"
HTML_OUT = ROOT / "docs" / "AZURE_GENAI_BEGINNER_HANDBOOK.html"
DOCX_OUT = ROOT / "docs" / "AZURE_GENAI_BEGINNER_HANDBOOK.docx"

BLUE = "0878D1"
NAVY = "152033"
MUTED = "56647A"
LIGHT_AMBER = "FFF4DF"


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([run_properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\[([^]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|_([^_]+)_|\*([^*]+)\*)"
)


def add_inline(paragraph, text: str) -> None:
    """Add the small Markdown inline subset used by this handbook."""
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        if match.group(2) is not None:
            target = match.group(3)
            if not target.startswith(("http://", "https://")):
                target = target
            add_hyperlink(paragraph, match.group(2), target)
        elif match.group(4) is not None:
            paragraph.add_run(match.group(4)).bold = True
        elif match.group(5) is not None:
            run = paragraph.add_run(match.group(5))
            run.style = "Inline Code"
        else:
            paragraph.add_run(match.group(6) or match.group(7)).italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    heading_sizes = {1: 24, 2: 17, 3: 13}
    for level, size in heading_sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else NAVY)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if level < 3 else 10)
        style.paragraph_format.space_after = Pt(6)

    code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string("EAF3FC")
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(5)
    code.paragraph_format.space_after = Pt(8)

    inline = document.styles.add_style("Inline Code", WD_STYLE_TYPE.CHARACTER)
    inline.font.name = "Consolas"
    inline.font.size = Pt(9)
    inline.font.color.rgb = RGBColor.from_string("075AA4")

    quote = document.styles["Quote"]
    quote.font.name = "Aptos"
    quote.font.size = Pt(10)
    quote.font.color.rgb = RGBColor.from_string(NAVY)
    quote.paragraph_format.left_indent = Inches(0.25)
    quote.paragraph_format.right_indent = Inches(0.12)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = footer.add_run("Azure GenAI Deployment Handbook for Beginners  •  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def markdown_to_docx(source: str) -> None:
    document = Document()
    configure_docx(document)

    lines = source.splitlines()
    section_titles = [
        line.removeprefix("## ")
        for line in lines
        if line.startswith("## ") and not line.startswith("### ")
    ]
    paragraph_buffer: list[str] = []
    in_code = False
    code_lines: list[str] = []
    seen_title = False

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(line.strip() for line in paragraph_buffer).strip()
        paragraph = document.add_paragraph()
        add_inline(paragraph, text)
        paragraph_buffer = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_paragraph()
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                set_paragraph_shading(paragraph, "091A2E")
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            flush_paragraph()
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2)
            paragraph = document.add_heading(level=level)
            add_inline(paragraph, title)
            if level == 1 and not seen_title:
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                subtitle = document.add_paragraph(
                    "Plain-language concepts, three application walkthroughs, Azure service decisions, CI/CD, cost, security, and cleanup"
                )
                subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
                subtitle.runs[0].font.size = Pt(11)
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                contents = document.add_paragraph()
                contents.add_run("Contents").bold = True
                contents.runs[0].font.size = Pt(20)
                contents.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
                document.add_paragraph(
                    "Use this map for reading order. Word's Navigation Pane also uses the heading structure to jump directly to any section."
                )
                for section_title in section_titles:
                    paragraph = document.add_paragraph(style="List Bullet")
                    paragraph.add_run(section_title)
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                seen_title = True
            continue

        if line.startswith("> "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Quote")
            add_inline(paragraph, line[2:])
            set_paragraph_shading(paragraph, LIGHT_AMBER)
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if bullet or numbered:
            flush_paragraph()
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet else "List Number"
            )
            add_inline(paragraph, (bullet or numbered).group(1))
            continue

        paragraph_buffer.append(line)

    flush_paragraph()

    properties = document.core_properties
    properties.title = "Azure GenAI Deployment Handbook for Beginners"
    properties.subject = "Plain-language Azure GenAI architecture and deployment learning guide"
    properties.author = "Testing_end_to_end learning project"
    properties.keywords = "Azure, GenAI, FastAPI, React, deployment, beginner, CI/CD"
    document.save(DOCX_OUT)


def markdown_to_html(source: str) -> None:
    md = markdown.Markdown(
        extensions=["fenced_code", "tables", "toc", "sane_lists"],
        extension_configs={"toc": {"permalink": False}},
    )
    body = md.convert(source)
    toc = md.toc
    html = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Azure GenAI Deployment Handbook for Beginners</title>
  <style>
    :root{{--ink:#152033;--muted:#56647a;--line:#d9e3ef;--blue:#0878d1;--navy:#071f3d;--paper:#f5f8fc;--card:#fff;--amber:#b46400}}
    *{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;background:var(--paper);color:var(--ink);font:17px/1.72 Inter,Segoe UI,Arial,sans-serif}}
    header{{padding:4rem 1.2rem;color:#fff;background:linear-gradient(125deg,#071f3d,#075aa4 55%,#5636ad)}}header>div{{max-width:1120px;margin:auto}}header h1{{font-size:clamp(2.2rem,5vw,4.3rem);line-height:1.04;margin:.4rem 0 1rem;max-width:980px}}header p{{max-width:820px;font-size:1.15rem}}
    .layout{{max-width:1380px;margin:auto;display:grid;grid-template-columns:300px minmax(0,1fr);gap:2rem;padding:2rem 1rem 5rem}}nav{{position:sticky;top:1rem;align-self:start;max-height:calc(100vh - 2rem);overflow:auto;background:#fff;border:1px solid var(--line);border-radius:1rem;padding:1rem}}nav strong{{display:block;color:var(--blue);margin-bottom:.7rem}}nav ul{{padding-left:1.15rem;margin:.25rem 0}}nav li{{margin:.3rem 0}}nav a{{color:#344760;text-decoration:none;font-size:.86rem}}nav a:hover{{color:var(--blue);text-decoration:underline}}
    article{{min-width:0;background:#fff;border:1px solid var(--line);border-radius:1.1rem;padding:clamp(1.1rem,3vw,3rem);box-shadow:0 12px 36px #193e6812}}h1,h2,h3{{line-height:1.2;scroll-margin-top:1rem}}article>h1:first-child{{display:none}}h2{{font-size:clamp(1.65rem,3vw,2.25rem);color:var(--blue);margin-top:2.8rem;border-top:1px solid var(--line);padding-top:1.7rem}}h2:first-of-type{{margin-top:0;border:0;padding-top:0}}h3{{font-size:1.28rem;margin-top:1.9rem}}p{{max-width:86ch}}blockquote{{margin:1.2rem 0;padding:1rem 1.2rem;border-left:5px solid var(--amber);background:#fff4df;border-radius:.65rem}}blockquote p{{margin:.2rem 0}}pre{{overflow:auto;background:#091a2e;color:#d9ecff;padding:1.1rem;border-radius:.8rem;font:14px/1.48 Consolas,monospace}}code{{background:#eef2f7;color:#075aa4;padding:.08rem .28rem;border-radius:.3rem}}pre code{{background:transparent;color:inherit;padding:0}}li{{margin:.36rem 0}}a{{color:#075aa4}}hr{{border:0;border-top:1px solid var(--line)}}
    .format-links{{display:flex;gap:.6rem;flex-wrap:wrap;margin-top:1.3rem}}.format-links a{{color:#fff;border:1px solid #ffffff66;border-radius:999px;padding:.42rem .75rem;text-decoration:none}}
    footer{{text-align:center;color:var(--muted);padding:2rem}}
    @media(max-width:940px){{.layout{{grid-template-columns:1fr}}nav{{position:static;max-height:none}}article{{padding:1.2rem}}}}
    @media print{{header{{background:#fff;color:var(--ink);padding:1rem}}nav{{display:none}}.layout{{display:block;padding:0}}article{{border:0;box-shadow:none}}}}
  </style>
</head>
<body>
  <header><div><div style="letter-spacing:.13em;text-transform:uppercase;font-weight:800;font-size:.78rem;opacity:.82">Plain-language learning edition</div><h1>Azure GenAI Deployment Handbook for Beginners</h1><p>Understand every request path, decision, YES/NO scenario, Azure service, cost, latency, security boundary, failure, and cleanup step before automating it.</p><div class="format-links"><a href="AZURE_GENAI_BEGINNER_HANDBOOK.md">Markdown source</a><a href="AZURE_GENAI_BEGINNER_HANDBOOK.docx">Download Word document</a><a href="GENAI_DEPLOYMENT_GUIDE.html">App-by-app guide</a></div></div></header>
  <div class="layout"><nav><strong>Contents</strong>{toc}</nav><article>{body}</article></div>
  <footer>Generated from AZURE_GENAI_BEGINNER_HANDBOOK.md · Keep Markdown, HTML, and Word synchronized.</footer>
</body>
</html>
"""
    HTML_OUT.write_text(html, encoding="utf-8")


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    markdown_to_html(source)
    markdown_to_docx(source)
    print(f"saved {HTML_OUT.relative_to(ROOT)}")
    print(f"saved {DOCX_OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
